"""Screen: EUROCONTROL's published DAILY all-cause pre-departure delay per airport against the v46 development stack (2025 months only).

The workbook (All_Pre-Departure_Delay.xlsx, sheet DATA: FLT_DATE, APT_ICAO, FLT_DEP_IFR_2, DLY_ALL_PRE_2 in minutes) is an aggregate of
the airport operators' off-block times, the clock the ranking file hides. Descriptive; nothing is selected or exported.

  1. what the figure is: DLY_ALL_PRE_2 against the day's sum of max(off-block - schedule, 0) from our own 2025 clocks
  2. what it says about our errors: delta = published total - the total the stack implies (off-block = takeoff - prediction), against the
     true discrepancy and the day's summed label error
  3. uses, each with the fitting side on the OTHER three development months:
       spread    delta spread over the day's rows in proportion to error variance (class x schedule-interval band), hinge linearised
       learned   residual LightGBM with day-level discrepancy features, against the same model without the published total
  4. identifiability: for every row with |error| > 1800 s that the total sees, the number of rows on its airport-day able to absorb the
     same discrepancy in the same direction

  --extract   workbook -> data/external/eurocontrol_aiu/All_Pre-Departure_Delay.parquet (ten airports, 2025 on)
  (default)   -> docs/clean_line_screens/daily_totals.json
"""
from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))
sys.path.insert(0, str(ROOT / "tools"))
from prc2026.artifacts import read_json, sha256_file, utc_now, write_json  # noqa: E402
from prc2026.features import ID  # noqa: E402
from prc2026_final import network_member, share_gate  # noqa: E402
from run_adsb_v11 import rmse  # noqa: E402
from run_callsign_fix import offsets, scope_table  # noqa: E402
from run_queue_trust import offsets_all, segments  # noqa: E402
from run_share_gate_reread import MONTHS, ROWS, STRETCH, full_stack  # noqa: E402

DIR = ROOT / "data/external/eurocontrol_aiu"
WORKBOOK, TABLE = DIR / "All_Pre-Departure_Delay.xlsx", DIR / "All_Pre-Departure_Delay.parquet"
AIRPORTS = ["EDDF", "EDDM", "EGLL", "EHAM", "LEBL", "LEMD", "LFPG", "LIRF", "LSZH", "LTFM"]
BANDS = [-1e6, -0.25, 0.25, 0.5, 1, 2, 4, 8, 12, 24, 1e6]
SEED = 20260929
OUT = ROOT / "docs/clean_line_screens/daily_totals.json"


def extract() -> None:
    from openpyxl import load_workbook
    rows = load_workbook(WORKBOOK, read_only=True, data_only=True)["DATA"].iter_rows(values_only=True)
    header = next(rows)
    cols = [i for i, h in enumerate(header) if h is not None]
    apt, year = header.index("APT_ICAO"), header.index("YEAR")
    keep = [[r[i] for i in cols] for r in rows if r[apt] in AIRPORTS and r[year] and int(r[year]) >= 2025]
    pd.DataFrame(keep, columns=[str(header[i]) for i in cols]).to_parquet(TABLE, index=False)
    print("kept", len(keep), "rows ->", TABLE.relative_to(ROOT).as_posix())


def published() -> pd.DataFrame:
    """Per (airport, day): the published minutes in seconds (pub) and the number of departures behind them (n_pub)."""
    p = pd.read_parquet(TABLE)
    p["day"] = pd.to_datetime(p["FLT_DATE"]).dt.strftime("%Y-%m-%d")
    p = p.set_index(["APT_ICAO", "day"]).rename_axis(["airport", "bday"])
    return pd.DataFrame({"pub": p["DLY_ALL_PRE_2"].astype(float) * 60, "n_pub": p["FLT_DEP_IFR_2"].astype(float)})


def stack_rows() -> pd.DataFrame:
    """Development rows with the v46 stack's value and the raw clocks."""
    rows = pd.read_parquet(ROOT / ROWS)
    cs = scope_table(read_json(ROOT / "configs/callsign_fix_v1.json")).set_index(ID)
    off_cs, off, seg = offsets(), offsets_all(), segments(MONTHS)
    k = read_json(ROOT / STRETCH)["candidates"]["two_sided_900"]["k_all_months"]
    gate_spec = read_json(ROOT / "configs/lirf_gate_v1.json")["candidates"]["u600_m600"]
    frames = []
    for month in MONTHS:
        f = rows[rows["month"] == month].reset_index(drop=True)
        mix, _ = share_gate.mix(f["committee"].to_numpy(dtype=float), f["nets"].to_numpy(dtype=float), network_member.V46_SHARE, network_member.V46_GATE)
        frames.append(f[[ID, "month", "airport", "label", "matched", "lirf_m", "lirf_u"]].assign(stack=full_stack(f, mix, cs, off_cs, seg, off, k, gate_spec)))
    t = pd.concat(frames, ignore_index=True)
    raw = pd.concat([pd.read_parquet(p, columns=[ID, "PHASE_mvt", "MVT_TIME_UTC_mvt", "BLOCK_TIME_UTC_mvt", "SCHED_TIME_UTC_mvt"])
                     for m in MONTHS for p in sorted((ROOT / "data/raw").glob(f"training_{m}-*.parquet"))], ignore_index=True)
    raw = raw[raw["PHASE_mvt"].eq("DEP")].drop_duplicates(ID)
    t = t.merge(raw.drop(columns="PHASE_mvt"), on=ID, how="left", validate="one_to_one")
    clocks = {c: pd.to_datetime(t[c], utc=True) for c in ["MVT_TIME_UTC_mvt", "BLOCK_TIME_UTC_mvt", "SCHED_TIME_UTC_mvt"]}
    t["a"] = (clocks["MVT_TIME_UTC_mvt"] - clocks["SCHED_TIME_UTC_mvt"]).dt.total_seconds()
    t["block_day"] = clocks["BLOCK_TIME_UTC_mvt"].dt.strftime("%Y-%m-%d")
    t["bday"] = (clocks["MVT_TIME_UTC_mvt"] - pd.to_timedelta(t["stack"], unit="s")).dt.strftime("%Y-%m-%d")  # label-free off-block day
    t["pos_true"], t["pos_hat"] = np.maximum(t["a"] - t["label"], 0.0), np.maximum(t["a"] - t["stack"], 0.0)
    t["r"] = t["label"] - t["stack"]
    t["se"] = t["r"] ** 2
    lirf = t["airport"].eq("LIRF")
    t["cls"] = np.where(lirf & ~t["lirf_m"] & ~t["lirf_u"], "LIRF_pop", np.where(t["matched"], t["airport"] + "_m", t["airport"] + "_u"))
    t["band"] = pd.cut(t["a"] / 3600, BANDS).astype(str)
    return t.drop(columns=["MVT_TIME_UTC_mvt", "BLOCK_TIME_UTC_mvt", "SCHED_TIME_UTC_mvt"])


def robust_sd(v: pd.Series) -> float:
    return 1.4826 * float(np.median(np.abs(v - np.median(v))))


def month_changes(t: pd.DataFrame, column: str) -> dict:
    out = {}
    for month, g in t.groupby("month"):
        y, o = g["label"].to_numpy(dtype=float), ((g["label"] > 0) & (g["label"] <= 7200)).to_numpy()
        before, after = g["stack"].to_numpy(dtype=float), g[column].to_numpy(dtype=float)
        out[month] = {"stack_rmse": rmse(y, before), "rmse": rmse(y, after), "change_pct": 100 * (rmse(y, after) / rmse(y, before) - 1),
                      "ordinary_rows_change_pct": 100 * (rmse(y[o], after[o]) / rmse(y[o], before[o]) - 1)}
    return out


def spread(t: pd.DataFrame, accuracy: pd.DataFrame) -> np.ndarray:
    from scipy.stats import norm
    new = t["stack"].to_numpy(dtype=float).copy()
    for month in MONTHS:
        fit, use = t[t["month"].ne(month)], t[t["month"].eq(month)]
        table, by_cls = fit.groupby(["cls", "band"])["se"].agg(["mean", "size"]), fit.groupby("cls")["se"].mean()
        key = pd.MultiIndex.from_frame(use[["cls", "band"]])
        v = np.where(table["size"].reindex(key).fillna(0).to_numpy() >= 20, table["mean"].reindex(key).to_numpy(), by_cls.reindex(use["cls"]).to_numpy())
        v = np.maximum(v, 900.0)
        sigma = accuracy[accuracy["month"].ne(month)].groupby(level="airport")["difference"].apply(robust_sd)
        y, a = use["stack"].to_numpy(dtype=float), use["a"].to_numpy(dtype=float)
        z = (a - y) / np.sqrt(v)
        g = pd.DataFrame({"airport": use["airport"].to_numpy(), "bday": use["bday"].to_numpy(), "e": (a - y) * norm.cdf(z) + np.sqrt(v) * norm.pdf(z), "h": v * norm.cdf(z) ** 2,
                          "gain": v * norm.cdf(z), "pub": use["pub"].to_numpy()})
        grp = g.groupby(["airport", "bday"])
        delta = g["pub"] - grp["e"].transform("sum")
        denom = grp["h"].transform("sum") + g["airport"].map(sigma).to_numpy() ** 2
        ok = (np.isfinite(delta) & use["complete"].to_numpy()).to_numpy()
        new[use.index.to_numpy()] = np.maximum(y - np.where(ok, g["gain"] * delta / denom, 0.0), -12.0)
    return new


def learned(t: pd.DataFrame) -> tuple[np.ndarray, np.ndarray]:
    import lightgbm as lgb
    f = t[["stack", "a", "pos_hat"]].copy()
    keys = [t["airport"], t["bday"]]
    f["cls_code"], f["apt_code"] = t["cls"].astype("category").cat.codes, t["airport"].astype("category").cat.codes
    f["n_day"] = t.groupby(["airport", "bday"])["a"].transform("size")
    implied = t.groupby(["airport", "bday"])["pos_hat"].transform("sum")
    for h in [1, 2, 4, 8]:
        f[f"n_a{h}"] = (t["a"] > 3600 * h).groupby(keys).transform("sum").astype(float)
        f[f"stake_a{h}"] = t["pos_hat"].where(t["a"] > 3600 * h, 0.0).groupby(keys).transform("sum")
    f["big_pred"] = t["stack"].where(t["stack"] > 3600, 0.0).groupby(keys).transform("sum")
    f["n_big_pred"] = (t["stack"] > 3600).groupby(keys).transform("sum").astype(float)
    f["rank_stake"] = t.groupby(["airport", "bday"])["pos_hat"].rank(ascending=False, method="first")
    f["rank_pred"] = t.groupby(["airport", "bday"])["stack"].rank(ascending=False, method="first")
    f["share_stake"] = t["pos_hat"] / implied.clip(lower=1.0)
    row_features = list(f.columns)
    f["delta"] = t["pub"] - implied
    f["delta_per_row"] = f["delta"] / f["n_day"]
    f["room_up"], f["room_down"] = t["pos_hat"] + f["delta"], t["stack"] - f["delta"]
    params = {"objective": "regression", "learning_rate": 0.05, "num_leaves": 31, "min_data_in_leaf": 200, "lambda_l2": 20.0, "feature_fraction": 0.8, "bagging_fraction": 0.8,
              "bagging_freq": 1, "verbose": -1, "seed": SEED, "num_threads": 8, "deterministic": True}
    ok = t["complete"].to_numpy()
    out = []
    for features in (row_features, list(f.columns)):
        new = t["stack"].to_numpy(dtype=float).copy()
        for month in MONTHS:
            fit, use = ok & t["month"].ne(month).to_numpy(), ok & t["month"].eq(month).to_numpy()
            model = lgb.train(params, lgb.Dataset(f.loc[fit, features], t.loc[fit, "r"], categorical_feature=["cls_code", "apt_code"]), num_boost_round=300)
            new[use] = np.maximum(new[use] + model.predict(f.loc[use, features]), -12.0)
        out.append(new)
    return out[0], out[1]


def identifiability(t: pd.DataFrame) -> dict:
    e = (t["pos_true"] - t["pos_hat"]).to_numpy()
    big = (t["r"].abs() > 1800).to_numpy()
    visible = np.abs(e) > 0.5 * t["r"].abs().to_numpy()
    total = float(t["se"].sum())
    groups = {k: (g["pos_hat"].to_numpy(), g["stack"].to_numpy()) for k, g in t.groupby(["airport", "bday"])}
    rows = []
    for i in np.flatnonzero(big & visible):
        pos_hat, stack = groups[(t["airport"].iat[i], t["bday"].iat[i])]
        rivals = int((pos_hat >= 0.8 * -e[i]).sum()) if e[i] < 0 else int((stack >= 0.8 * e[i]).sum())
        rows.append({"direction": "label_above_prediction" if e[i] < 0 else "label_below_prediction", "rivals": rivals, "se": float(t["se"].iat[i]), "airport": t["airport"].iat[i]})
    q = pd.DataFrame(rows)
    q["band"] = pd.cut(q["rivals"], [0, 3, 10, 1e6], labels=["1-3", "4-10", "over 10"])
    table = {f"{d}|{b}": {"rows": int(len(g)), "share_of_all_sse_pct": 100 * float(g["se"].sum()) / total} for (d, b), g in q.groupby(["direction", "band"], observed=True)}
    return {"rows_with_error_over_1800s": int(big.sum()), "sse_share_pct": 100 * float(t.loc[big, "se"].sum()) / total,
            "seen_by_the_total": {"rows": int((big & visible).sum()), "sse_share_pct": 100 * float(t.loc[big & visible, "se"].sum()) / total},
            "by_direction_and_rivals": table, "airports_of_rows_with_1_to_3_rivals": q[q["rivals"] <= 3].groupby("airport").size().to_dict()}


def main() -> None:
    if "--extract" in sys.argv:
        extract()
        return
    t = stack_rows()
    pub = published()
    t = t.join(pub, on=["airport", "bday"])
    # an airport-day is read only when our rows cover the departures behind the published figure
    t["complete"] = (t["pub"].notna() & (t.groupby(["airport", "bday"])["a"].transform("size") >= 0.9 * t["n_pub"])).to_numpy()
    report = {"created_at_utc": utc_now(), "audit_used": False, "official_scores_used": "none", "months": MONTHS,
              "source": {"workbook": "https://www.eurocontrol.int/performance/data/download/xls/All_Pre-Departure_Delay.xlsx", "sha256": sha256_file(WORKBOOK),
                         "release": "2026-09-16, period to 2026-07-31", "terms": "EUROCONTROL performance data disclaimer: source to be mentioned, non-commercial, no modification without permission; not an open licence",
                         "nature": "daily aggregate of airport-operator off-block times minus schedule (positive part), i.e. of the clock hidden on the ranking rows"},
              "control_month_rmse": {m: rmse(g["label"].to_numpy(dtype=float), g["stack"].to_numpy(dtype=float)) for m, g in t.groupby("month")}}

    own = t.groupby(["airport", "block_day"]).agg(true=("pos_true", "sum"), rows=("a", "size"), month=("month", "first")).rename_axis(["airport", "bday"]).join(pub, how="inner").dropna()
    own = own[own["rows"] >= 0.9 * own["n_pub"]]
    own["difference"] = own["pub"] - own["true"]
    report["published_against_own_clocks_minutes"] = {a: {"airport_days": int(len(g)), "median_abs": float(g["difference"].abs().median() / 60), "p90_abs": float(g["difference"].abs().quantile(0.9) / 60),
                                                          "robust_sd": robust_sd(g["difference"]) / 60, "correlation": float(g["pub"].corr(g["true"]))} for a, g in own.groupby(level="airport")}
    day = t[t["complete"]].groupby(["airport", "bday"]).agg(true=("pos_true", "sum"), hat=("pos_hat", "sum"), r=("r", "sum"), pub=("pub", "first")).dropna()
    report["discrepancy"] = {a: {"sd_true_discrepancy_minutes": float((g["true"] - g["hat"]).std() / 60), "correlation_published_with_true": float((g["pub"] - g["hat"]).corr(g["true"] - g["hat"])),
                                 "correlation_with_minus_summed_label_error": float((g["pub"] - g["hat"]).corr(-g["r"]))} for a, g in day.groupby(level="airport")}
    t["spread"] = spread(t, own)
    t["learned_control"], t["learned"] = learned(t)
    report["airport_days_read"] = int(t[t["complete"]].groupby(["airport", "bday"]).ngroups)
    report["rows_outside_a_read_airport_day"] = int((~t["complete"]).sum())
    report["uses"] = {k: month_changes(t, k) for k in ["spread", "learned_control", "learned"]}
    report["uses"]["learned_minus_its_control_pct"] = {m: report["uses"]["learned"][m]["change_pct"] - report["uses"]["learned_control"][m]["change_pct"] for m in MONTHS}
    report["identifiability"] = identifiability(t)
    report["code"] = {"tools/screen_daily_totals.py": sha256_file(ROOT / "tools/screen_daily_totals.py")}
    write_json(OUT, report)
    for a, e in report["published_against_own_clocks_minutes"].items():
        print(a, {k: round(v, 2) for k, v in e.items()}, {k: round(v, 2) for k, v in report["discrepancy"][a].items()})
    for k, e in report["uses"].items():
        print(k, {m: (round(v["change_pct"], 3) if isinstance(v, dict) else round(v, 3)) for m, v in e.items()})
    print(report["identifiability"])


if __name__ == "__main__":
    main()
